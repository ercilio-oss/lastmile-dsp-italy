#!/usr/bin/env python3
"""
QuickSight Dashboard Data Downloader
=====================================
Automates login to AWS QuickSight and downloads data from specified
dashboard sheets in Excel format by scraping the rendered pivot table DOM.

Usage:
    python quicksight_downloader.py                  # Download all sheets
    python quicksight_downloader.py --sheet 0        # Download first sheet only
    python quicksight_downloader.py --headless       # Run without browser UI
    python quicksight_downloader.py --config my.json # Use custom config

Requirements:
    pip install playwright openpyxl
    playwright install chromium
"""

import argparse
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

try:
    from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeout
except ImportError:
    print("ERROR: Playwright not installed. Run:")
    print("  pip install playwright")
    print("  playwright install chromium")
    sys.exit(1)

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
except ImportError:
    print("ERROR: openpyxl not installed. Run:")
    print("  pip install openpyxl")
    sys.exit(1)


# ─── Configuration ───────────────────────────────────────────────────────────

def load_config(config_path="config.json"):
    """Load configuration from JSON file."""
    config_file = Path(__file__).parent / config_path
    if not config_file.exists():
        print(f"ERROR: Config file not found: {config_file}")
        print("Copy config.json.example to config.json and fill in your credentials.")
        sys.exit(1)

    with open(config_file) as f:
        config = json.load(f)

    # Allow environment variable overrides for security
    config["quicksight"]["account_name"] = os.environ.get(
        "QS_ACCOUNT", config["quicksight"].get("account_name", "agl-iris-qs-ext")
    )
    config["quicksight"]["username"] = os.environ.get(
        "QS_USERNAME", config["quicksight"]["username"]
    )
    config["quicksight"]["password"] = os.environ.get(
        "QS_PASSWORD", config["quicksight"]["password"]
    )

    if config["quicksight"]["password"] == "CHANGE_ME_DO_NOT_STORE_REAL_PASSWORD":
        password = os.environ.get("QS_PASSWORD", "")
        if not password:
            print("ERROR: Password not configured.")
            print("Either:")
            print("  1. Set QS_PASSWORD environment variable, or")
            print("  2. Update password in config.json (less secure)")
            sys.exit(1)
        config["quicksight"]["password"] = password

    return config


# ─── Logging ─────────────────────────────────────────────────────────────────

def log(message, level="INFO"):
    """Simple timestamped logging."""
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{timestamp}] [{level}] {message}")


# ─── Core Automation ─────────────────────────────────────────────────────────

class QuickSightDownloader:
    """Handles QuickSight login and data download via browser automation."""

    # Column names for the pivot table rows, ordered by hierarchy depth
    ROW_COLUMNS = [
        "region_id",
        "country",
        "provider_company_short_code",
        "delivery_station_code",
        "service_type_name",
        "year_week",
        "reporting_date",
        "Route Length",
        "Route Rate",
    ]

    def __init__(self, config, headless=False):
        self.config = config
        self.headless = headless
        self.download_dir = self._setup_download_dir()
        self.playwright = None
        self.browser = None
        self.page = None

    def _setup_download_dir(self):
        """Create timestamped download directory."""
        base_dir = Path(__file__).parent / self.config["download"]["output_dir"]
        timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
        download_dir = base_dir / timestamp
        download_dir.mkdir(parents=True, exist_ok=True)
        log(f"Download directory: {download_dir}")
        return download_dir

    def start_browser(self):
        """Launch browser with download handling."""
        self.playwright = sync_playwright().start()
        self.browser = self.playwright.chromium.launch(
            headless=self.headless,
            args=["--no-sandbox", "--disable-dev-shm-usage"],
        )
        self.context = self.browser.new_context(
            accept_downloads=True,
            viewport={"width": 1920, "height": 1080},
        )
        self.page = self.context.new_page()
        self.page.set_default_timeout(60000)  # 60s timeout
        log("Browser started")

    def stop_browser(self):
        """Clean up browser resources."""
        if self.browser:
            self.browser.close()
        if self.playwright:
            self.playwright.stop()
        log("Browser closed")

    def _click_submit(self):
        """Click the Next / Sign in submit button."""
        btn = self.page.wait_for_selector(
            'button[type="submit"]', timeout=5000
        )
        btn.click()
        time.sleep(3)

    def _dismiss_modals(self):
        """
        Dismiss any modal overlays (welcome modal, announcements, etc.)
        that QuickSight shows on first visit.
        """
        time.sleep(2)
        dismissed = False

        # Strategy 1: Click close/dismiss/OK buttons inside the modal
        modal_close_selectors = [
            '[data-automation-id="welcome-modal"] button',
            '[data-automation-id="welcome-modal"] [aria-label="Close"]',
            '.MuiDialog-root button[aria-label="Close"]',
            '.MuiDialog-root button[aria-label="close"]',
            '.MuiDialog-root button:has-text("Close")',
            '.MuiDialog-root button:has-text("OK")',
            '.MuiDialog-root button:has-text("Got it")',
            '.MuiDialog-root button:has-text("Dismiss")',
            '.MuiDialog-root button:has-text("Skip")',
            '.MuiDialog-root button:has-text("Continue")',
            '.MuiDialog-root [class*="close"]',
            '.MuiDialog-root [class*="Close"]',
            '[data-testid="close-button"]',
            '[data-testid="dismiss-button"]',
        ]

        for selector in modal_close_selectors:
            try:
                btn = self.page.wait_for_selector(selector, timeout=2000)
                if btn and btn.is_visible():
                    btn.click(force=True)
                    log(f"Dismissed modal via: {selector}")
                    dismissed = True
                    time.sleep(1)
                    break
            except (PlaywrightTimeout, Exception):
                continue

        # Strategy 2: Press Escape
        if not dismissed:
            try:
                modal = self.page.query_selector(
                    '[data-automation-id="welcome-modal"], .MuiDialog-root'
                )
                if modal and modal.is_visible():
                    self.page.keyboard.press("Escape")
                    log("Dismissed modal via Escape key")
                    dismissed = True
                    time.sleep(1)
            except Exception:
                pass

        # Strategy 3: Click backdrop
        if not dismissed:
            try:
                backdrop = self.page.query_selector(
                    '.MuiBackdrop-root, [class*="backdrop"], [class*="Backdrop"]'
                )
                if backdrop and backdrop.is_visible():
                    backdrop.click(force=True)
                    log("Dismissed modal by clicking backdrop")
                    dismissed = True
                    time.sleep(1)
            except Exception:
                pass

        # Strategy 4: Force-remove from DOM
        if not dismissed:
            try:
                removed = self.page.evaluate("""
                    () => {
                        const modals = document.querySelectorAll(
                            '[data-automation-id="welcome-modal"], .MuiDialog-root'
                        );
                        if (modals.length > 0) {
                            modals.forEach(m => m.remove());
                            return true;
                        }
                        return false;
                    }
                """)
                if removed:
                    log("Dismissed modal by removing from DOM")
                    dismissed = True
                    time.sleep(1)
            except Exception:
                pass

        if not dismissed:
            log("No modal detected (or already dismissed)")

        return dismissed

    def login(self):
        """
        Log into QuickSight.
        Handles 2-step (username → password) or 3-step (account → username → password).
        """
        qs = self.config["quicksight"]
        account_name = qs.get("account_name", "agl-iris-qs-ext")

        login_url = (
            "https://us-east-1.quicksight.aws.amazon.com/sn/auth/signin"
            "?enable-sso=0"
            "&redirect_uri=https%3A%2F%2Fus-east-1.quicksight.aws.amazon.com"
            "%2Fsn%2Fstart%3Fdirectory-alias%3D" + account_name +
            "%26enable-sso%3D0%26state%3DhashArgs%2523%26isauthcode%3Dtrue"
        )

        log(f"Navigating to QuickSight login (account: {account_name})...")
        self.page.goto(login_url, wait_until="networkidle")
        time.sleep(3)

        # Check if we need account name step
        page_text = self.page.content()

        if f"Sign in to" not in page_text or account_name not in page_text:
            log(f"Account name step detected. Entering: {account_name}")
            try:
                account_field = self.page.wait_for_selector(
                    'input[type="text"]', timeout=10000
                )
                account_field.fill(account_name)
                self._click_submit()
                self.page.screenshot(
                    path=str(self.download_dir / "login_step0_account.png")
                )
            except PlaywrightTimeout:
                log("Could not find account field", "ERROR")
                raise

        # Step 1: Username
        log(f"Step 1/2: Entering username: {qs['username']}")
        try:
            username_field = self.page.wait_for_selector(
                'input[type="text"]', timeout=10000
            )
            username_field.fill(qs["username"])
            self._click_submit()
            self.page.screenshot(
                path=str(self.download_dir / "login_step1_username.png")
            )
            log("Username submitted, waiting for password page...")
        except PlaywrightTimeout:
            self.page.screenshot(
                path=str(self.download_dir / "login_step1_error.png")
            )
            raise Exception("Login failed at username step. See screenshot.")

        # Step 2: Password
        log("Step 2/2: Entering password...")
        try:
            password_field = self.page.wait_for_selector(
                'input[type="password"]', timeout=15000
            )
            password_field.fill(qs["password"])
            log("Password entered, clicking Sign in...")
            self._click_submit()
            self.page.screenshot(
                path=str(self.download_dir / "login_step2_signin.png")
            )
        except PlaywrightTimeout:
            self.page.screenshot(
                path=str(self.download_dir / "login_step2_error.png")
            )
            raise Exception("Login failed at password step. See screenshot.")

        # Verify login success
        log("Waiting for post-login redirect...")
        time.sleep(5)

        current_url = self.page.url
        success_indicators = ["/start", "/dashboards", "/account/", "/analyses"]
        fail_indicators = ["signin", "/auth/"]

        if any(ind in current_url for ind in success_indicators):
            log("Login successful!")
        elif any(ind in current_url.lower() for ind in fail_indicators):
            self.page.screenshot(
                path=str(self.download_dir / "login_error_final.png")
            )
            log(f"Still on login page: {current_url}", "ERROR")
            raise Exception(
                "Login failed. Check account name, username, and password."
            )
        else:
            log(f"Redirected to: {current_url} — assuming login OK")

        time.sleep(2)
        self._dismiss_modals()

    # ─── Data Extraction (DOM Scraping) ──────────────────────────────────────

    def download_sheet_data(self, sheet_url, sheet_name):
        """
        Navigate to a dashboard sheet and extract all pivot table data
        by scraping the rendered DOM, then save as Excel.
        """
        log(f"Navigating to sheet: {sheet_name}")
        self.page.goto(sheet_url, wait_until="networkidle")
        time.sleep(8)  # Let dashboard visuals render fully

        # Dismiss any modal
        self._dismiss_modals()
        time.sleep(2)

        # Take a screenshot for debugging
        self.page.screenshot(
            path=str(self.download_dir / f"{sheet_name}_loaded.png")
        )
        log(f"Screenshot saved for {sheet_name}")

        # Detect all pivot tables / data tables on the page
        tables_data = self._extract_all_tables()

        if not tables_data:
            log(f"No table data found on {sheet_name}", "WARN")
            return

        # Save each table to a separate sheet in one Excel file
        self._save_to_excel(tables_data, sheet_name)

    def _extract_all_tables(self):
        """
        Extract data from all pivot tables on the current page.
        Returns a list of dicts: [{name, columns, rows}, ...]
        """
        tables = []

        # Find all pivot table visuals by looking for the visual title + row headers
        result = self.page.evaluate("""
            () => {
                const tables = [];

                // Find visual containers that contain pivot tables
                // Pivot tables have elements with class "row-headers-container"
                const rowContainers = document.querySelectorAll('.row-headers-container');

                rowContainers.forEach((container, tableIdx) => {
                    // Walk up to find the visual's title
                    let titleText = 'Table_' + (tableIdx + 1);
                    let parent = container;
                    for (let i = 0; i < 15 && parent; i++) {
                        const titleEl = parent.querySelector(
                            '.text-box-visual-container .ql-content, ' +
                            '[class*="visual-title"], ' +
                            '.dashboard-visual-title'
                        );
                        if (titleEl) {
                            titleText = titleEl.textContent.trim();
                            break;
                        }
                        parent = parent.parentElement;
                    }

                    // Get the column header names from the pivot
                    // Row-level headers are: region_id, country, etc. (visible as column buttons)
                    const colButtons = container.closest('[class*="sn-pivot"]') ||
                                       container.closest('[class*="visual"]') ||
                                       container.parentElement?.parentElement;

                    // --- Extract Row Headers ---
                    // Row headers have data-automation-id="sn-table-row-X.X.X..."
                    const rowEls = container.querySelectorAll(
                        '[data-automation-id^="sn-table-row-"]'
                    );

                    // Build hierarchy: group by path depth
                    const rowHierarchy = [];
                    rowEls.forEach(el => {
                        const aid = el.getAttribute('data-automation-id');
                        if (aid === 'sn-table-row-header-selection-indicator') return;
                        const title = el.querySelector('.title');
                        const text = title ? title.textContent.trim() : '';
                        if (!text) return;

                        const path = aid.replace('sn-table-row-', '');
                        const depth = path.split('.').length;
                        rowHierarchy.push({ path, depth, text });
                    });

                    // --- Extract Value Cells ---
                    // Values are in div.printable-cell inside div.cell.numeric
                    // They're positioned with CSS (left, top) in a grid
                    const gridContainer = container.parentElement?.querySelector('.grid') ||
                                          container.closest('[class*="sn-pivot"]')?.querySelector('.grid');

                    let valueCells = [];
                    if (gridContainer) {
                        const cells = gridContainer.querySelectorAll('.cell.numeric .printable-cell');
                        cells.forEach(c => {
                            const style = c.parentElement.getAttribute('style') || '';
                            const leftMatch = style.match(/left:\\s*(\\d+)px/);
                            const topMatch = style.match(/top:\\s*(\\d+)px/);
                            valueCells.push({
                                text: c.textContent.trim(),
                                left: leftMatch ? parseInt(leftMatch[1]) : 0,
                                top: topMatch ? parseInt(topMatch[1]) : 0
                            });
                        });
                    }

                    // --- Get Value Column Headers ---
                    // These are "Routes Accepted Total" and "Total Cost" (from the column header area)
                    let valueColumnNames = [];
                    const columnHeaderArea = container.parentElement?.querySelector(
                        '[class*="column-headers"], [class*="column-header-row"]'
                    );
                    if (columnHeaderArea) {
                        const headers = columnHeaderArea.querySelectorAll('.printable-cell, .title');
                        headers.forEach(h => {
                            const txt = h.textContent.trim();
                            if (txt) valueColumnNames.push(txt);
                        });
                    }

                    // If we couldn't find column headers from the DOM, try the button approach
                    if (valueColumnNames.length === 0) {
                        // Look for column header buttons nearby
                        const visual = container.closest('[class*="visual"]') ||
                                       container.parentElement?.parentElement?.parentElement;
                        if (visual) {
                            const btns = visual.querySelectorAll('button');
                            btns.forEach(b => {
                                const txt = b.textContent.trim();
                                if (txt === 'Routes Accepted Total' || txt === 'Total Cost') {
                                    valueColumnNames.push(txt);
                                }
                            });
                        }
                    }

                    tables.push({
                        name: titleText,
                        rowHierarchy,
                        valueCells,
                        valueColumnNames
                    });
                });

                return tables;
            }
        """)

        if not result:
            log("No pivot table data found in DOM", "WARN")
            return []

        # Now flatten the hierarchical data into tabular rows
        all_tables = []
        for table_raw in result:
            flat_rows = self._flatten_pivot_data(table_raw)
            if flat_rows:
                all_tables.append({
                    "name": table_raw["name"],
                    "value_columns": table_raw.get("valueColumnNames", []),
                    "rows": flat_rows,
                })

        return all_tables

    def _flatten_pivot_data(self, table_raw):
        """
        Flatten the hierarchical pivot table row data into flat rows.

        The row hierarchy has elements like:
          path="0.0.0.0.0.0.0.0.0" depth=9 text="55.18"  (Route Rate)
          path="0.0.0.0.0.0.0.0"   depth=8 text="120"    (Route Length)
          path="0.0.0.0.0.0.0"     depth=7 text="Mar 7, 2026" (reporting_date)
          ...
          path="0"                  depth=1 text="2"       (region_id)

        For each leaf row (max depth), we collect all ancestors to build
        the full flat row.
        """
        hierarchy = table_raw.get("rowHierarchy", [])
        value_cells = table_raw.get("valueCells", [])
        value_columns = table_raw.get("valueColumnNames", [])

        if not hierarchy:
            return []

        # Determine the max depth (leaf level)
        max_depth = max(h["depth"] for h in hierarchy)

        # Build a lookup: path -> text
        path_to_text = {}
        for h in hierarchy:
            path_to_text[h["path"]] = h["text"]

        # Identify leaf rows (those at max_depth)
        leaf_rows = [h for h in hierarchy if h["depth"] == max_depth]

        # Determine the number of value columns per row
        num_value_cols = len(value_columns) if value_columns else 2

        # Sort value cells by top position to match row order, then by left
        sorted_values = sorted(value_cells, key=lambda v: (v["top"], v["left"]))

        # Group value cells by row (same top position)
        value_rows = []
        current_top = None
        current_row = []
        for vc in sorted_values:
            if current_top is None or vc["top"] != current_top:
                if current_row:
                    value_rows.append(current_row)
                current_row = [vc["text"]]
                current_top = vc["top"]
            else:
                current_row.append(vc["text"])
        if current_row:
            value_rows.append(current_row)

        # Build flat rows by combining hierarchy path + values
        flat_rows = []
        for i, leaf in enumerate(leaf_rows):
            # Reconstruct all ancestor values from the path
            parts = leaf["path"].split(".")
            row_values = []

            for depth_idx in range(len(parts)):
                ancestor_path = ".".join(parts[: depth_idx + 1])
                text = path_to_text.get(ancestor_path, "")
                row_values.append(text)

            # Add value cells for this row
            if i < len(value_rows):
                row_values.extend(value_rows[i])

            flat_rows.append(row_values)

        return flat_rows

    def _save_to_excel(self, tables_data, sheet_name):
        """Save extracted table data to a styled Excel file."""
        filename = f"{sheet_name}_{datetime.now().strftime('%Y-%m-%d_%H%M%S')}.xlsx"
        filepath = self.download_dir / filename

        wb = openpyxl.Workbook()
        # Remove the default sheet
        wb.remove(wb.active)

        # Styling
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for table in tables_data:
            # Clean sheet name (Excel max 31 chars)
            ws_name = re.sub(r'[\\/*?:\[\]]', '_', table["name"])[:31]
            ws = wb.create_sheet(title=ws_name)

            rows = table["rows"]
            value_columns = table.get("value_columns", [])

            if not rows:
                ws.cell(row=1, column=1, value="No data found")
                continue

            # Determine column headers
            # Row columns from hierarchy + value columns
            max_row_depth = max(len(r) - len(value_columns) for r in rows) if rows else 0
            # Use known column names if depth matches
            if max_row_depth <= len(self.ROW_COLUMNS):
                col_headers = self.ROW_COLUMNS[:max_row_depth]
            else:
                col_headers = [f"Column_{i+1}" for i in range(max_row_depth)]

            col_headers.extend(value_columns if value_columns else
                             [f"Value_{i+1}" for i in range(len(rows[0]) - max_row_depth)])

            # Write headers
            for col_idx, header in enumerate(col_headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            # Write data rows
            for row_idx, row_data in enumerate(rows, 2):
                for col_idx, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border

                    # Try to convert to number
                    cleaned = str(value).replace(",", "")
                    try:
                        num_val = float(cleaned)
                        cell.value = num_val
                        cell.number_format = '#,##0.##'
                    except (ValueError, TypeError):
                        cell.value = value

            # Auto-width columns
            for col_idx in range(1, len(col_headers) + 1):
                max_width = len(str(col_headers[col_idx - 1]))
                for row_idx in range(2, min(len(rows) + 2, 100)):
                    cell_value = ws.cell(row=row_idx, column=col_idx).value
                    if cell_value:
                        max_width = max(max_width, len(str(cell_value)))
                ws.column_dimensions[
                    openpyxl.utils.get_column_letter(col_idx)
                ].width = min(max_width + 2, 35)

            # Freeze header row
            ws.freeze_panes = "A2"

            # Add auto-filter
            ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(col_headers))}1"

            log(f"  Table '{table['name']}': {len(rows)} rows x {len(col_headers)} columns")

        wb.save(str(filepath))
        log(f"Saved Excel file: {filepath}")
        return filepath

    # ─── Scroll & Load All Data ──────────────────────────────────────────────

    def _scroll_to_load_all_data(self):
        """
        Scroll the sheet view to ensure all data tables are loaded.
        QuickSight lazy-loads visuals that are off-screen.
        """
        log("Scrolling page to load all visuals...")

        self.page.evaluate("""
            () => {
                const scrollable = document.querySelector(
                    '.scrollable-sheet-view, [class*="sheet-view"][class*="scrollable"]'
                );
                if (scrollable) {
                    // Scroll to bottom to trigger lazy loading
                    scrollable.scrollTop = scrollable.scrollHeight;
                }
            }
        """)
        time.sleep(3)

        # Scroll back to top
        self.page.evaluate("""
            () => {
                const scrollable = document.querySelector(
                    '.scrollable-sheet-view, [class*="sheet-view"][class*="scrollable"]'
                );
                if (scrollable) {
                    scrollable.scrollTop = 0;
                }
            }
        """)
        time.sleep(2)

    def run(self, sheet_indices=None):
        """Execute the full download workflow."""
        sheets = self.config["quicksight"]["dashboard_sheets"]

        if sheet_indices is not None:
            sheets = [sheets[i] for i in sheet_indices if i < len(sheets)]

        log(f"Starting download for {len(sheets)} sheet(s)...")

        try:
            self.start_browser()
            self.login()

            for i, sheet in enumerate(sheets):
                log(f"--- Sheet {i + 1}/{len(sheets)}: {sheet['name']} ---")

                # Navigate and wait for content
                self.page.goto(sheet["url"], wait_until="networkidle")
                time.sleep(8)

                # Dismiss modals
                self._dismiss_modals()
                time.sleep(2)

                # Scroll to load all visuals
                self._scroll_to_load_all_data()

                # Take debug screenshot
                self.page.screenshot(
                    path=str(self.download_dir / f"{sheet['name']}_loaded.png")
                )
                log(f"Screenshot saved for {sheet['name']}")

                # Extract and save data
                tables_data = self._extract_all_tables()

                if tables_data:
                    self._save_to_excel(tables_data, sheet["name"])
                else:
                    log(f"No table data found on {sheet['name']}", "WARN")

                time.sleep(2)

            # List downloaded files
            files = list(self.download_dir.iterdir())
            data_files = [f for f in files if f.name.endswith((".xlsx", ".csv"))]
            log(f"Download complete! {len(data_files)} data file(s) saved to {self.download_dir}")

            if not data_files:
                log(
                    "No data files were downloaded. Screenshots saved for debugging.",
                    "WARN",
                )

        except Exception as e:
            log(f"Error: {e}", "ERROR")
            try:
                self.page.screenshot(
                    path=str(self.download_dir / "error_screenshot.png")
                )
            except Exception:
                pass
            raise
        finally:
            self.stop_browser()

        return self.download_dir


# ─── CLI ─────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Download data from AWS QuickSight dashboards"
    )
    parser.add_argument(
        "--config",
        default="config.json",
        help="Path to config file (default: config.json)",
    )
    parser.add_argument(
        "--sheet",
        type=int,
        nargs="+",
        help="Sheet index(es) to download (0-based). Omit for all.",
    )
    parser.add_argument(
        "--headless",
        action="store_true",
        help="Run browser in headless mode (no UI)",
    )
    args = parser.parse_args()

    config = load_config(args.config)
    downloader = QuickSightDownloader(config, headless=args.headless)
    download_dir = downloader.run(sheet_indices=args.sheet)
    print(f"\nFiles saved to: {download_dir}")


if __name__ == "__main__":
    main()
